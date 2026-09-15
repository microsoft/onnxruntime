#!/usr/bin/env python3
"""Generate the QDQ MatMul matrix, run it on an NPU, and write XLSX results."""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


@dataclass(frozen=True)
class Category:
    title: str
    slug: str
    quantization: str
    qdq_profile: str


@dataclass(frozen=True)
class Variant:
    slug: str
    signedness: str
    bit_width: int
    symmetry: str
    omit_zero_point: bool = False


@dataclass(frozen=True)
class Shape:
    label: str
    slug: str
    input_shape: tuple[int, ...]
    weight_shape: tuple[int, int]


CATEGORIES = (
    Category(
        "PER-TENSOR-WTS / com.microsoft DQ",
        "per_tensor_microsoft",
        "per-tensor",
        "microsoft",
    ),
    Category(
        "PER-TENSOR-WTS / ONNX DQ",
        "per_tensor_onnx",
        "per-tensor",
        "onnx",
    ),
    Category(
        "PER-CHANNEL-WTS / com.microsoft DQ",
        "per_channel_microsoft",
        "per-channel",
        "microsoft",
    ),
    Category(
        "PER-CHANNEL-WTS / ONNX DQ",
        "per_channel_onnx",
        "per-channel",
        "onnx",
    ),
    Category(
        "BLOCKWISE WTS / ONNX DQ",
        "blockwise_onnx",
        "blockwise",
        "onnx",
    ),
)

VARIANTS = (
    Variant("int4_symmetric_no_zp", "signed", 4, "symmetric", True),
    Variant("int4_symmetric_zero_zp", "signed", 4, "symmetric"),
    Variant("int4_asymmetric", "signed", 4, "asymmetric"),
    Variant("int8_symmetric_no_zp", "signed", 8, "symmetric", True),
    Variant("int8_symmetric_zero_zp", "signed", 8, "symmetric"),
    Variant("int8_asymmetric", "signed", 8, "asymmetric"),
    Variant("uint4_symmetric", "unsigned", 4, "symmetric"),
    Variant("uint4_asymmetric", "unsigned", 4, "asymmetric"),
    Variant("uint8_symmetric", "unsigned", 8, "symmetric"),
    Variant("uint8_asymmetric", "unsigned", 8, "asymmetric"),
)

SHAPES = (
    Shape("[1, 768], [768, 512]", "1x768_768x512", (1, 768), (768, 512)),
    Shape(
        "[1, 2520, 768], [768, 768]",
        "1x2520x768_768x768",
        (1, 2520, 768),
        (768, 768),
    ),
)

QUANTIZATION_HEADERS = (
    ("int4-symmetric", 2),
    ("int4-asym", 1),
    ("int8-symmetric", 2),
    ("int8-asym", 1),
    ("uint4-symmetric", 1),
    ("uint4-asym", 1),
    ("uint8-symmmetric", 1),
    ("uint8-asym", 1),
)

ZERO_POINT_HEADERS = (
    "No zp",
    "All 0 zp",
    "zp",
    "No zp",
    "All 0 zp",
    "zp",
    "zp",
    "zp",
    "zp",
    "zp",
)

PLACEMENT_PATTERN = re.compile(
    r"(?:All nodes|Node\(s\)) placed on \[([^\]]+)\]\. "
    r"Number of nodes: (\d+)"
)
PLACED_NODE_PATTERN = re.compile(r"\]\s+(\S+)\s+\([^)]*\)\s*$")
MAE_PATTERN = re.compile(r"^- Mean abs error:\s*(\S+)\s*$", re.MULTILINE)
COSINE_PATTERN = re.compile(
    r"^- Cosine similarity:\s*(\S+)\s*$", re.MULTILINE
)


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=script_dir / "unit-models" / "qdq-matmul-test-suite",
        help="Generated model directory.",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=script_dir / "qdq_matmul_npu_results.xlsx",
        help="Output XLSX path.",
    )
    parser.add_argument(
        "--logs-dir",
        type=Path,
        default=script_dir / "qdq-matmul-test-suite-logs",
        help="Directory for run_acc.py logs.",
    )
    parser.add_argument(
        "--provider",
        choices=("vitisai", "qnn", "openvino"),
        default="vitisai",
        help="NPU provider passed to run_acc.py.",
    )
    parser.add_argument(
        "--provider-option",
        action="append",
        default=[],
        metavar="KEY=VALUE",
        help="Provider option passed to run_acc.py; may be repeated.",
    )
    parser.add_argument(
        "--block-size",
        type=int,
        choices=(32, 128),
        default=32,
        help="Block size for blockwise models.",
    )
    parser.add_argument(
        "--block-axis",
        type=int,
        choices=(0, 1),
        default=0,
        help="Blocked axis for blockwise models.",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=0,
        help="Model and run_acc.py input seed.",
    )
    return parser.parse_args()


def create_workbook() -> tuple[Workbook, Worksheet, dict[tuple[int, int], int]]:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Results"
    data_rows: dict[tuple[int, int], int] = {}
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    subheader_fill = PatternFill("solid", fgColor="EAF2F8")
    table_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for category_index, category in enumerate(CATEGORIES):
        header_row = category_index * 5 + 1
        zp_row = header_row + 1
        worksheet.cell(header_row, 1, category.title)
        worksheet.cell(zp_row, 1, "ZP present or not")

        column = 2
        for label, width in QUANTIZATION_HEADERS:
            worksheet.cell(header_row, column, label)
            if width > 1:
                worksheet.merge_cells(
                    start_row=header_row,
                    start_column=column,
                    end_row=header_row,
                    end_column=column + width - 1,
                )
            column += width

        for variant_index, label in enumerate(ZERO_POINT_HEADERS):
            worksheet.cell(zp_row, variant_index + 2, label)

        for shape_index, shape in enumerate(SHAPES):
            row = header_row + 2 + shape_index
            data_rows[(category_index, shape_index)] = row
            worksheet.cell(row, 1, shape.label)
            worksheet.row_dimensions[row].height = 48

        for cell in worksheet[header_row]:
            cell.fill = header_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for cell in worksheet[zp_row]:
            cell.fill = subheader_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in worksheet.iter_rows(
            min_row=header_row,
            max_row=header_row + 3,
            min_col=1,
            max_col=11,
        ):
            for cell in row:
                cell.border = table_border

    worksheet.column_dimensions["A"].width = 36
    for column in range(2, 12):
        worksheet.column_dimensions[get_column_letter(column)].width = 29
    worksheet.freeze_panes = "B3"
    return workbook, worksheet, data_rows


def run_command(command: list[str]) -> tuple[int, str]:
    completed = subprocess.run(
        command,
        check=False,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    return completed.returncode, completed.stdout.replace("\x00", "")


def generate_model(
    python: str,
    generator: Path,
    output: Path,
    category: Category,
    variant: Variant,
    shape: Shape,
    block_size: int,
    block_axis: int,
    seed: int,
) -> tuple[int, str]:
    command = [
        python,
        str(generator),
        "--output",
        str(output),
        "--input-shape",
        *(str(value) for value in shape.input_shape),
        "--weight-shape",
        *(str(value) for value in shape.weight_shape),
        "--weight-quantization",
        category.quantization,
        "--weight-signedness",
        variant.signedness,
        "--weight-bit-width",
        str(variant.bit_width),
        "--weight-symmetry",
        variant.symmetry,
        "--qdq-profile",
        category.qdq_profile,
        "--seed",
        str(seed),
        "--add-vitisai-metadata"
    ]
    if variant.omit_zero_point:
        command.append("--omit-weight-zero-point")
    if category.quantization == "blockwise":
        command.extend(
            ["--block-size", str(block_size), "--block-axis", str(block_axis)]
        )
    return run_command(command)


def run_accuracy(
    python: str,
    runner: Path,
    model: Path,
    provider: str,
    provider_options: list[str],
    seed: int,
) -> tuple[int, str]:
    command = [
        python,
        str(runner),
        str(model),
        "--provider",
        provider,
        "--seed",
        str(seed),
        "--log-severity-level",
        "0"
        # ,"--warmup-iterations",
        # "1000"
    ]
    for option in provider_options:
        command.extend(["--provider-option", option])
    return run_command(command)


def extract_result(log: str) -> tuple[int, int, tuple[str, ...], str, str]:
    placement_start = log.rfind("Node placements")
    if placement_start < 0:
        raise ValueError("NPU run log does not contain 'Node placements'")

    placement_log = log[placement_start:]
    placements = PLACEMENT_PATTERN.findall(placement_log)
    if not placements:
        raise ValueError("NPU node placement counts were not found")

    cpu_ops = sum(
        int(count)
        for provider, count in placements
        if provider == "CPUExecutionProvider"
    )
    npu_ops = sum(
        int(count)
        for provider, count in placements
        if provider != "CPUExecutionProvider"
    )
    cpu_op_names: list[str] = []
    current_provider = ""
    remaining_nodes = 0
    for line in placement_log.splitlines():
        placement_match = PLACEMENT_PATTERN.search(line)
        if placement_match:
            current_provider = placement_match.group(1)
            remaining_nodes = int(placement_match.group(2))
            continue

        if remaining_nodes == 0:
            continue

        node_match = PLACED_NODE_PATTERN.search(line)
        if node_match:
            if current_provider == "CPUExecutionProvider":
                cpu_op_names.append(node_match.group(1))
            remaining_nodes -= 1

    mae_matches = MAE_PATTERN.findall(log)
    cosine_matches = COSINE_PATTERN.findall(log)
    if not mae_matches or not cosine_matches:
        raise ValueError("accuracy metrics were not found")
    return cpu_ops, npu_ops, tuple(cpu_op_names), mae_matches[-1], cosine_matches[-1]


def format_result(
    cpu_ops: int,
    npu_ops: int,
    cpu_op_names: tuple[str, ...],
    mae: str,
    cosine: str,
) -> str:
    cpu_op_summary = (
        f" ({', '.join(cpu_op_names)})" if cpu_op_names else ""
    )
    return (
        f"CPU ops: {cpu_ops}{cpu_op_summary}, NPU ops: {npu_ops}\n"
        f"MAE vs CPU run = {mae}\n"
        f"Cosine similarity vs CPU run = {cosine}"
    )


def format_error(message: str) -> str:
    one_line_message = " ".join(message.split())
    return (
        "CPU ops: ERROR, NPU ops: ERROR\n"
        f"MAE vs CPU run = ERROR: {one_line_message}\n"
        "Cosine similarity vs CPU run = ERROR"
    )


def final_error(output: str, fallback: str) -> str:
    lines = [line.strip() for line in output.splitlines() if line.strip()]
    return lines[-1] if lines else fallback


def iter_cases() -> Iterator[
    tuple[int, Category, int, Shape, int, Variant]
]:
    for category_index, category in enumerate(CATEGORIES):
        for shape_index, shape in enumerate(SHAPES):
            for variant_index, variant in enumerate(VARIANTS):
                yield (
                    category_index,
                    category,
                    shape_index,
                    shape,
                    variant_index,
                    variant,
                )


def main() -> None:
    args = parse_args()
    script_dir = Path(__file__).resolve().parent
    generator = script_dir / "generate_qdq_matmul_model.py"
    runner = script_dir / "run_acc.py"
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.logs_dir.mkdir(parents=True, exist_ok=True)
    args.workbook.parent.mkdir(parents=True, exist_ok=True)

    workbook, worksheet, data_rows = create_workbook()
    workbook.save(args.workbook)
    total = len(CATEGORIES) * len(SHAPES) * len(VARIANTS)
    failed_models: set[Path] = set()

    for completed, case in enumerate(iter_cases(), start=1):
        category_index, category, shape_index, shape, variant_index, variant = (
            case
        )
        stem = f"{category.slug}_{shape.slug}_{variant.slug}"
        model_path = args.output_dir / f"{stem}.onnx"
        log_path = args.logs_dir / f"{stem}.log"
        print(f"[generate {completed}/{total}] {stem}", flush=True)
        return_code, output = generate_model(
            sys.executable,
            generator,
            model_path,
            category,
            variant,
            shape,
            args.block_size,
            args.block_axis,
            args.seed,
        )
        if return_code == 0:
            continue

        failed_models.add(model_path)
        log_path.write_text(output, encoding="utf-8")
        row = data_rows[(category_index, shape_index)]
        cell = worksheet.cell(
            row,
            variant_index + 2,
            format_error(final_error(output, "model generation failed")),
        )
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        workbook.save(args.workbook)

    for completed, case in enumerate(iter_cases(), start=1):
        category_index, category, shape_index, shape, variant_index, variant = (
            case
        )
        stem = f"{category.slug}_{shape.slug}_{variant.slug}"
        model_path = args.output_dir / f"{stem}.onnx"
        if model_path in failed_models:
            continue

        log_path = args.logs_dir / f"{stem}.log"
        print(f"[run {completed}/{total}] {stem}", flush=True)
        return_code, output = run_accuracy(
            sys.executable,
            runner,
            model_path,
            args.provider,
            args.provider_option,
            args.seed,
        )
        log_path.write_text(output, encoding="utf-8")
        if return_code != 0:
            result = format_error(final_error(output, "run_acc.py failed"))
            passed = False
        else:
            try:
                extracted_result = extract_result(output)
                result = format_result(*extracted_result)
                cpu_ops, _, _, _, cosine = extracted_result
                passed = cpu_ops == 0 and float(cosine) == 1.0
            except ValueError as error:
                result = format_error(str(error))
                passed = False

        row = data_rows[(category_index, shape_index)]
        cell = worksheet.cell(row, variant_index + 2, result)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if passed:
            cell.fill = PatternFill("solid", fgColor="C6EFCE")
        workbook.save(args.workbook)

    print(f"Saved results to {args.workbook.resolve()}")


if __name__ == "__main__":
    main()
